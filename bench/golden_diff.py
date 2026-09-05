"""Golden diff: partitioned and streamed evaluation must equal a whole-file run.

`eval_grid` loads and evaluates the workbook in one piece, which is the
behaviour this library has to preserve exactly. Every other entry point is
checked against it cell by cell, on generated workbooks that exercise both the
partitioned path and the whole-file fallback, plus optional fixture or downloaded files.

Run from the package root: python bench/golden_diff.py
"""

import glob
import io
import os
import sys

import openpyxl

import formualizer_partitioned

# min_formulas=0: partition whenever it is safe, so the sweep exercises the
# partitioned path rather than the size heuristic that governs production use.

# Optional real workbooks to check when present under this package's fixtures
# or the downloaded bench/data/fed corpus.
REAL_FILES = sorted(
    glob.glob(os.path.join(os.path.dirname(__file__), "fixtures", "*.xlsx"))
    + glob.glob(os.path.join(os.path.dirname(__file__), "data", "fed", "*.xlsx"))
)


def workbook_bytes(build):
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row_local(wb):
    """Independent per-row formulas: the shape that partitions best."""
    ws = wb.active
    ws.title = "S"
    for r in range(1, 61):
        ws.cell(r, 1, r)
        ws.cell(r, 2, f"n{r}")
        ws.cell(r, 3, f"=A{r}*2")
        ws.cell(r, 4, f'=IF(A{r}>30,"hi",B{r})')


def cross_sheet(wb):
    ws, other = wb.active, wb.create_sheet("T")
    ws.title = "S"
    for r in range(1, 61):
        ws.cell(r, 1, r)
        ws.cell(r, 2, f"=A{r}+1")
        other.cell(r, 1, f"=S!B{r}*3")


def chain(wb):
    """A running chain cannot be split, so this must fall back."""
    ws = wb.active
    ws.title = "S"
    ws.cell(1, 1, 1)
    for r in range(2, 61):
        ws.cell(r, 1, f"=A{r - 1}+1")


def dynamic_refs(wb):
    """INDIRECT hides its precedents, so this must fall back."""
    ws = wb.active
    ws.title = "S"
    for r in range(1, 61):
        ws.cell(r, 1, r)
        ws.cell(r, 2, f'=INDIRECT("A"&{r})+1')


def mixed_types(wb):
    ws = wb.active
    ws.title = "S"
    import datetime

    for r in range(1, 41):
        ws.cell(r, 1, r * 1.5)
        ws.cell(r, 2, datetime.date(2024, 1, 1) + datetime.timedelta(days=r))
        ws.cell(r, 3, r % 2 == 0)
        ws.cell(r, 4, f"=A{r}*2")
        ws.cell(r, 5, f"=C{r}")
        ws.cell(r, 6, "=1/0")


def whole_column_aggregates(wb):
    """Every row divides by an aggregate over a whole data column.

    Without the prelude each formula reads the whole column, so every bounding
    box covers the sheet and the file falls back. The prelude folds each
    aggregate to a number, and the result must still equal a whole-file run.

    The row count is above one prelude chunk, so the fold has to add chunk
    totals rather than read the column once.
    """
    ws = wb.active
    ws.title = "S"
    rows = 1300
    for r in range(1, rows + 1):
        ws.cell(r, 1, r)
        ws.cell(r, 2, (r % 7) - 3)
        ws.cell(r, 3, ["EMEA", "APAC", "AMER"][r % 3])
        ws.cell(r, 4, f"=A{r}/SUM(A:A)*100")
        ws.cell(r, 5, f"=A{r}-AVERAGE(A:A)")
        ws.cell(r, 6, f"=IF(A{r}>MAX(B:B),1,0)")
        ws.cell(r, 7, f"=A{r}+MIN(B:B)")
        ws.cell(r, 8, f'=A{r}/COUNTIF(C:C,"EMEA")')
        ws.cell(r, 9, f'=A{r}+SUMIF(C:C,"APAC",A:A)')
        ws.cell(r, 10, f"=COUNT(A:A)-A{r}")


def unfoldable_aggregate(wb):
    """`MEDIAN` needs every value at once, so the prelude must leave it alone
    and the file must keep the behaviour it had before."""
    ws = wb.active
    ws.title = "S"
    for r in range(1, 61):
        ws.cell(r, 1, r % 13)
        ws.cell(r, 2, f"=A{r}-MEDIAN(A:A)")


def constant_name(wb):
    """A name defined as a constant, which the whole-file loader drops.

    The loader keeps only range definitions, so a whole-file run answers
    `#NAME?` here. The reader therefore refuses the file instead of resolving
    the constant, which would make the answer depend on the layout that ran.
    `constant_names_still_unresolved` watches that assumption.
    """
    from openpyxl.workbook.defined_name import DefinedName

    ws = wb.active
    ws.title = "S"
    wb.defined_names.add(DefinedName("_Order1", attr_text="0"))
    for r in range(1, 61):
        ws.cell(r, 1, r)
        ws.cell(r, 2, f"=A{r}+_Order1")


def constant_names_still_unresolved(data):
    """Fail when the whole-file loader starts resolving constant names.

    While it drops them, resolving a constant name in a mini-workbook would
    give a different answer from the whole-file run this library must
    reproduce, so `graph::parse_defined_target` refuses such a name and
    `partition_verdict` sends the file to the whole-file path.

    When this check fails the loader has started to keep constant names.
    Resolve them in `parse_defined_target` (a literal target beside the range
    target), define them per batch in `partition::define_static_name`, and drop
    the `unsupported_refs` count they carry today.
    """
    cell = formualizer_partitioned.eval_grid(data)["S"][0][1]
    resolved = not (isinstance(cell, dict) and cell.get("kind") == "Name")
    print(f"  {'constant-name support':<28} "
          f"{'whole-file now resolves constants' if resolved else 'whole-file drops constants (as assumed)':<44}"
          f"{'':>6}       {'FAILURES' if resolved else 'ok'}")
    if resolved:
        return [f"whole-file run resolved a constant defined name ({cell!r}); "
                "re-enable constant names in the reader (see this function's docstring)"]
    return []


def compare(name, data):
    grid = formualizer_partitioned.eval_grid(data)
    plan = formualizer_partitioned.partition_plan(data, None, 0.9, 0)

    failures = []

    part = formualizer_partitioned.eval_partitioned(data, None, 0.9, 0)
    if part.keys() != grid.keys():
        failures.append(f"sheet set differs: {sorted(part)} vs {sorted(grid)}")
    for sheet in grid:
        if len(part.get(sheet, [])) != len(grid[sheet]):
            failures.append(f"{sheet}: row count differs")
            continue
        for i, (want, got) in enumerate(zip(grid[sheet], part[sheet])):
            if want != got:
                failures.append(f"eval_partitioned {sheet}!row{i + 1}: {want!r} != {got!r}")

    streamed = 0
    for sheet, rn, row in formualizer_partitioned.eval_rows(data, None, 0.9, 0):
        streamed += 1
        want = grid[sheet][rn - 1]
        if row != want:
            failures.append(f"eval_rows {sheet}!row{rn}: {want!r} != {row!r}")
    expected_rows = sum(len(rows) for rows in grid.values())
    if streamed != expected_rows:
        failures.append(f"streamed {streamed} rows, grid has {expected_rows}")

    mode = "partitioned" if plan["partitioned"] else f"fallback ({plan['fallback_reason']})"
    folded = plan.get("prelude_folded", 0)
    note = f" +{folded} folded" if folded else ""
    status = "ok" if not failures else f"{len(failures)} FAILURES"
    print(f"  {name:<28} {mode + note:<44} {expected_rows:>6} rows  {status}")
    for f in failures[:5]:
        print(f"      {f}")
    return failures


def main():
    cases = [
        ("row-local", row_local),
        ("cross-sheet", cross_sheet),
        ("chain", chain),
        ("dynamic refs", dynamic_refs),
        ("mixed types", mixed_types),
        ("whole-column aggregates", whole_column_aggregates),
        ("unfoldable aggregate", unfoldable_aggregate),
        ("constant name", constant_name),
    ]
    failures = []
    print("generated workbooks")
    for name, build in cases:
        failures += compare(name, workbook_bytes(build))
    failures += constant_names_still_unresolved(workbook_bytes(constant_name))

    real = [p for p in REAL_FILES if os.path.exists(p)]
    if real:
        print("real workbooks")
        for path in real:
            with open(path, "rb") as fh:
                failures += compare(os.path.basename(path)[:28], fh.read())

    print("FAIL" if failures else "PASS")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
