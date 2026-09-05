"""Run the golden diff across a corpus of real workbooks.

Generated fixtures are not enough: partitioning meets things real spreadsheets
do that a fixture author does not think of. This sweeps every xlsx under a
directory, compares `eval_partitioned` and `eval_rows`
against `eval_grid`, and reports which files disagree and why the rest were not
partitioned. For each disagreement, it also says whether the first differing
formula matches the value cached in the workbook by Excel.

Get a corpus first, for example SpreadsheetBench (912 real questions collected
from Excel forums, ~2700 input workbooks):

    curl -L -o ssb.tar.gz \\
      "https://huggingface.co/datasets/KAKA22/SpreadsheetBench/resolve/main/spreadsheetbench_912_v0.1.tar.gz?download=true"
    tar xzf ssb.tar.gz

Run from the package root: python bench/corpus_diff.py DIR [limit]
"""

import collections
import glob
import os
import re
import resource
import sys
import time
from io import BytesIO

import formualizer_partitioned

# min_formulas=0: partition whenever it is safe, so the sweep exercises the
# partitioned path rather than the size heuristic that governs production use.


def compare(data):
    """Return the differences, strategy, and evaluation times for one workbook."""
    timings = collections.Counter()
    # One instant for all three runs. Each call otherwise pins its own, so a
    # file calling NOW() would differ by a second whenever two calls straddle
    # a second boundary.
    now = time.time()
    started = time.perf_counter()
    base = formualizer_partitioned.eval_grid(data, now)
    timings["eval_grid"] += time.perf_counter() - started
    started = time.perf_counter()
    part = formualizer_partitioned.eval_partitioned(data, None, 0.9, 0, now=now)
    timings["eval_partitioned"] += time.perf_counter() - started
    diffs = 0
    first = None

    if base.keys() != part.keys():
        diffs += 1
        first = ("sheet set", sorted(base), sorted(part))
    for sheet in base.keys() & part.keys():
        if len(base[sheet]) != len(part[sheet]):
            diffs += 1
            first = first or (sheet, "row count", len(base[sheet]), len(part[sheet]))
            continue
        for ri, (rb, rp) in enumerate(zip(base[sheet], part[sheet])):
            if len(rb) != len(rp):
                diffs += 1
                first = first or (sheet, f"row {ri + 1} width", len(rb), len(rp))
                continue
            for ci, (cb, cp) in enumerate(zip(rb, rp)):
                if cb != cp:
                    diffs += 1
                    first = first or (sheet, f"r{ri + 1}c{ci + 1}", cb, cp)

    started = time.perf_counter()
    rows = formualizer_partitioned.eval_rows(data, None, 0.9, 0, now=now)
    seen_rows = collections.Counter()
    for sheet, rn, row in rows:
        seen_rows[sheet] = max(seen_rows[sheet], rn)
        if sheet not in base or rn > len(base[sheet]):
            diffs += 1
            first = first or ("unexpected row", sheet, rn)
            continue
        want = base[sheet][rn - 1]
        if row != want:
            diffs += 1
            if first is None:
                for ci, (expected, actual) in enumerate(zip(want, row), 1):
                    if expected != actual:
                        first = (sheet, f"r{rn}c{ci}", expected, actual)
                        break
    for sheet, expected in base.items():
        if expected and seen_rows[sheet] != len(expected):
            diffs += 1
            first = first or ("row count", sheet, len(expected), seen_rows[sheet])
    timings["eval_rows"] += time.perf_counter() - started
    return diffs, first, rows.strategy, timings


def cached_evidence(data, first):
    """Say whether Excel's cached value matches either result at the first diff."""
    if first is None or len(first) != 4:
        return "cached value unavailable for this difference"
    sheet, ref, whole, partitioned = first
    found = re.fullmatch(r"r(\d+)c(\d+)", ref)
    if not found:
        return "cached value unavailable for this difference"
    row, col = map(int, found.groups())
    import openpyxl
    formulas = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False)
    cached = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet not in formulas.sheetnames:
            return "cached value unavailable: sheet is missing"
        formula_cell = formulas[sheet].cell(row, col)
        if formula_cell.data_type != "f":
            return "difference is not in a formula cell"
        value = cached[sheet].cell(row, col).value
    finally:
        formulas.close()
        cached.close()
    matches = []
    if value == whole:
        matches.append("whole")
    if value == partitioned:
        matches.append("partitioned")
    return f"cached={value!r:.40} matches={'+'.join(matches) or 'neither'}"


def display_first(first):
    if first is None:
        return None
    return tuple(repr(value)[:40] if i >= 2 else value for i, value in enumerate(first))


def main():
    if len(sys.argv) < 2:
        print("usage: python bench/corpus_diff.py DIR [limit]", file=sys.stderr)
        return 2
    root = sys.argv[1]
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 0

    files = sorted(glob.glob(os.path.join(root, "**", "*.xlsx"), recursive=True))
    files = [f for f in files if "_answer" not in os.path.basename(f)]
    if limit:
        files = files[:limit]

    verdicts = collections.Counter()
    strategies = collections.Counter()
    errors = collections.Counter()
    bad = []
    checked = 0
    timings = collections.Counter()
    total_started = time.perf_counter()

    for path in files:
        started = time.perf_counter()
        with open(path, "rb") as fh:
            data = fh.read()
        timings["read"] += time.perf_counter() - started
        try:
            started = time.perf_counter()
            plan = formualizer_partitioned.partition_plan(data, None, 0.9, 0)
            timings["partition_plan"] += time.perf_counter() - started
        except Exception as exc:
            errors[f"plan: {type(exc).__name__}: {exc}"[:90]] += 1
            continue
        if not plan["partitioned"]:
            verdicts[plan["fallback_reason"]] += 1
            continue
        verdicts["PARTITION"] += 1
        try:
            diffs, first, strategy, measured = compare(data)
        except Exception as exc:
            errors[f"eval: {type(exc).__name__}: {exc}"[:90]] += 1
            continue
        timings.update(measured)
        strategies[strategy] += 1
        checked += 1
        if diffs:
            bad.append((os.path.relpath(path, root), diffs, display_first(first), cached_evidence(data, first)))

    print(f"{len(files)} workbooks under {root}")
    for reason, n in verdicts.most_common():
        print(f"   {n:>5}  {100 * n / max(len(files), 1):5.1f}%  {reason}")
    print(f"\ncompared {checked} partitioned workbooks; {len(bad)} disagree")
    for strategy, n in strategies.most_common():
        print(f"   {n:>5}  eval_rows used {strategy}")
    for name, diffs, first, cached in bad[:15]:
        print(f"   {name:<34} {diffs:>6} diffs  {first}; {cached}")
    if errors:
        print("\nerrors:")
        for msg, n in errors.most_common(10):
            print(f"   {n:>5}  {msg}")

    elapsed = time.perf_counter() - total_started
    peak_mb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024
    measured = sum(timings.values())
    print("\ntiming:")
    for name in ["read", "partition_plan", "eval_grid", "eval_partitioned", "eval_rows"]:
        print(f"   {timings[name]:>8.2f} s  {name}")
    print(f"   {elapsed - measured:>8.2f} s  comparisons and Python overhead")
    print(f"   {elapsed:>8.2f} s  total")
    print(f"   {peak_mb:>8.1f} MB peak RSS (the diff holds both full grids)")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
